import json
import openpyxl
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

with open('geocode_results.json', encoding='utf-8') as f:
    geo = json.load(f)
with open('viabilidade_resultado.json', encoding='utf-8') as f:
    viab = json.load(f)

SRC = '/sessions/affectionate-eloquent-bohr/mnt/uploads/LeadsFibraVendaGross.xlsx'
OUT = 'LeadsFibraVendaGross_viabilidade.xlsx'

wb = openpyxl.load_workbook(SRC)
ws = wb['Planilha1']

COL_VIAB = 28  # AB

dv = DataValidation(type='list', formula1='Planilha2!$A$1:$A$3', allow_blank=True, showDropDown=False)
ws.add_data_validation(dv)

n_fibra = n_hfc = n_nao = n_semdados = n_falhageo = 0
low_conf_types = {'APPROXIMATE', 'GEOMETRIC_CENTER'}

for row_idx in range(2, ws.max_row + 1):
    linha = str(row_idx)
    g = geo.get(linha, {})
    v = viab.get(linha, {})
    cell = ws.cell(row=row_idx, column=COL_VIAB)
    dv.add(cell)

    if g.get('status') != 'OK':
        n_falhageo += 1
        cell.value = None
        cell.comment = Comment(f"Não foi possível geocodificar o endereço (Google retornou: {g.get('status')}). Verifique CEP/número manualmente.", "Apex - verificação automática")
        continue

    motivo = v.get('motivo')
    if motivo == 'sem_kmz_na_cidade':
        n_semdados += 1
        cell.value = None
        cell.comment = Comment(f"Sem dado de cobertura (KMZ) para a cidade '{v.get('cidade_geo')}' no arquivo kmz_fixa.kmz enviado. Não é 'sem viabilidade' — é 'não verificado'.", "Apex - verificação automática")
        continue

    resultado = v.get('viabilidade')
    cell.value = resultado
    if resultado == 'Fibra':
        n_fibra += 1
    elif resultado == 'HFC':
        n_hfc += 1
    elif resultado == 'NÃO':
        n_nao += 1

    partes_comment = [f"Endereço geocodificado: {g.get('formatted')}"]
    dg = v.get('dist_gpon_m'); dh = v.get('dist_hfc_m')
    if dg is not None: partes_comment.append(f"Distância até polígono GPON mais próximo: {dg}m (limite 150m)")
    if dh is not None: partes_comment.append(f"Distância até polígono HFC mais próximo: {dh}m (limite 50m)")
    loc_type = g.get('location_type')
    if loc_type in low_conf_types:
        partes_comment.append(f"ATENÇÃO: precisão da geocodificação é '{loc_type}' (baixa/média confiança — pode não ser o endereço exato). Confira manualmente.")
    cell.comment = Comment("\n".join(partes_comment), "Apex - verificação automática")

wb.save(OUT)
print(f"Fibra={n_fibra} HFC={n_hfc} NÃO={n_nao} sem_dados_kmz={n_semdados} falha_geocodificacao={n_falhageo}")
print("total:", n_fibra+n_hfc+n_nao+n_semdados+n_falhageo)
